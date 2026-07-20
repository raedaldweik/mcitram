# NCGR meeting deck v2 — fewer words, real product screenshots.
# SAS corporate presentation style.
import os
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

NAVY   = "032954"
NAVY_2 = "0B3A6B"
BLUE   = "0766D1"
SKY    = "8FC1F2"
TINT   = "EDF4FC"
BORDER = "D7E4F4"
INK    = "1F2A37"
GRAY   = "53565A"
FAINT  = "9AA5B1"
WHITE  = "FFFFFF"
TEAL   = "0E7490"
GREEN  = "0B6E4F"
AMBER  = "B45309"

HERE   = Path(__file__).resolve().parent
ASSETS = str(HERE / "assets")
LOGO   = os.getenv("NCGR_LOGO", str(Path(__file__).resolve().parents[2]
                   / "frontend" / "public" / "ncgr-logo.png"))

W, H = 13.333, 7.5
MX = 0.62
CW = W - 2 * MX

prs = Presentation()
prs.slide_width = Inches(W)
prs.slide_height = Inches(H)
BLANK = prs.slide_layouts[6]

FOOT_L = "Company Confidential — For Internal Use Only"
FOOT_C = "Copyright © SAS Institute Inc. All rights reserved."


def slide():
    return prs.slides.add_slide(BLANK)

def box(s, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, 0)
    return tf

def para(tf, first=False, align=PP_ALIGN.LEFT, space_after=0, space_before=0,
         line=None):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_after = Pt(space_after)
    p.space_before = Pt(space_before)
    if line:
        p.line_spacing = line
    return p

def run(p, text, size, color, bold=False, italic=False, name="Arial"):
    r = p.add_run()
    r.text = text
    f = r.font
    f.size, f.bold, f.italic, f.name = Pt(size), bold, italic, name
    f.color.rgb = RGBColor.from_string(color)
    return r

def rect(s, x, y, w, h, fill, line_color=None, round_=False, adj=0.055):
    st = MSO_SHAPE.ROUNDED_RECTANGLE if round_ else MSO_SHAPE.RECTANGLE
    sp = s.shapes.add_shape(st, Inches(x), Inches(y), Inches(w), Inches(h))
    if round_:
        try:
            sp.adjustments[0] = adj
        except Exception:
            pass
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = RGBColor.from_string(fill)
    if line_color is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = RGBColor.from_string(line_color)
        sp.line.width = Pt(1)
    sp.shadow.inherit = False
    return sp

def footer(s, n, dark=False):
    c = "5E7A9B" if dark else FAINT
    tf = box(s, MX, H - 0.34, 5.6, 0.25)
    run(para(tf, True), FOOT_L, 7.5, c)
    tf = box(s, W / 2 - 2.4, H - 0.34, 4.8, 0.25)
    run(para(tf, True, align=PP_ALIGN.CENTER), FOOT_C, 7.5, c)
    tf = box(s, W - MX - 2.0, H - 0.34, 2.0, 0.25)
    run(para(tf, True, align=PP_ALIGN.RIGHT), f"sas.com   ·   {n}", 7.5, c)

def header(s, kicker, title, accent=BLUE, title_color=NAVY):
    rect(s, MX, 0.52, 0.42, 0.045, accent)
    tf = box(s, MX + 0.56, 0.40, CW - 0.56, 0.3)
    run(para(tf, True), kicker.upper(), 11, accent, bold=True)
    tf = box(s, MX, 0.74, CW, 0.62)
    run(para(tf, True), title, 27, title_color, bold=True)

def browser(s, x, y, w, img, url="ncgr-copilot.up.railway.app"):
    """A screenshot inside a light browser chrome. Images are 1920x1200."""
    bar = 0.34
    ih = w * (1200 / 1920)
    rect(s, x - 0.02, y - 0.02, w + 0.04, ih + bar + 0.04, WHITE,
         line_color=BORDER)
    rect(s, x, y, w, bar, "EDF1F7")
    for i, c in enumerate(("E8736F", "F2C14E", "6FBF73")):
        d = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + 0.14 + i * 0.17),
                               Inches(y + bar / 2 - 0.04),
                               Inches(0.08), Inches(0.08))
        d.fill.solid(); d.fill.fore_color.rgb = RGBColor.from_string(c)
        d.line.fill.background(); d.shadow.inherit = False
    pw = min(3.4, w - 1.1)
    rect(s, x + 0.72, y + 0.065, pw, bar - 0.13, WHITE, line_color=BORDER,
         round_=True, adj=0.5)
    tf = box(s, x + 0.88, y + 0.05, pw - 0.2, bar - 0.1,
             anchor=MSO_ANCHOR.MIDDLE)
    run(para(tf, True), url, 8.5, GRAY)
    s.shapes.add_picture(f"{ASSETS}/{img}", Inches(x), Inches(y + bar),
                         width=Inches(w))
    return ih + bar

def side_bullets(s, items, accent, y=1.75, size=15):
    tf = box(s, MX, y, 4.35, 4.6)
    for i, it in enumerate(items):
        p = para(tf, i == 0, space_after=16, line=1.15)
        run(p, "▪  ", size, accent, bold=True)
        run(p, it, size, INK)
    return tf


# ════════════════════════════════════════════════════════════════════
# 1 · TITLE (navy)
# ════════════════════════════════════════════════════════════════════
s = slide()
rect(s, 0, 0, W, H, NAVY)
rect(s, 0, 0, W, 0.07, BLUE)
s.shapes.add_picture(LOGO, Inches(W - MX - 1.05), Inches(0.55),
                     height=Inches(1.05))
rect(s, MX, 2.3, 0.6, 0.055, BLUE)
tf = box(s, MX, 2.56, 11.2, 0.35)
run(para(tf, True), "SAS MIDDLE EAST  ·  AGENTIC AI IN ACTION", 13, SKY,
    bold=True)
tf = box(s, MX, 3.0, 12.1, 1.05)
run(para(tf, True), "The NCGR AI Copilot", 46, WHITE, bold=True)
tf = box(s, MX, 4.12, 12.1, 0.55)
p = para(tf, True)
run(p, "«مُعين»", 22, SKY, bold=True)
run(p, "   —   your platform, your data, your choice of AI model.", 17, SKY)
tf = box(s, MX, 6.28, 9.0, 0.75)
p = para(tf, True, space_after=3)
run(p, "National Center for Government Resources Systems (NCGR)", 13, WHITE,
    bold=True)
p = para(tf)
run(p, "SAS Middle East  ·  July 2026", 12, "AFC2DA")
footer(s, 1, dark=True)

# ════════════════════════════════════════════════════════════════════
# 2 · THE VISION — three tiles, few words
# ════════════════════════════════════════════════════════════════════
s = slide()
header(s, "The Vision", "Your Copilot, On Your Terms")
cards = [
    ("Your AI model", "Cloud today. Fully on-premises tomorrow. "
                      "Swap it anytime."),
    ("Your SAS platform", "Real data, real models, real dashboards — "
                          "not chatbot guesses."),
    ("Your language", "English  ⇄  Arabic. One click."),
]
cw, gapx = 3.86, 0.26
for i, (t, d) in enumerate(cards):
    x = MX + i * (cw + gapx)
    rect(s, x, 2.0, cw, 3.4, TINT, line_color=BORDER, round_=True)
    rect(s, x, 2.0, cw, 0.09, BLUE)
    tf = box(s, x + 0.34, 2.42, cw - 0.68, 0.6)
    run(para(tf, True), f"{i+1}", 30, BLUE, bold=True)
    tf = box(s, x + 0.34, 3.18, cw - 0.68, 0.6)
    run(para(tf, True, line=1.05), t, 19, NAVY, bold=True)
    tf = box(s, x + 0.34, 3.85, cw - 0.68, 1.4)
    run(para(tf, True, line=1.25), d, 13.5, INK)
footer(s, 2)

# ════════════════════════════════════════════════════════════════════
# 3 · MEET YOUR AI TEAM — with the real dropdown
# ════════════════════════════════════════════════════════════════════
s = slide()
header(s, "One Chat · Four Agents", "Meet Your AI Team")
team = [
    ("SAS Viya Copilot", BLUE, "Data, models & dashboards on request"),
    ("Investigation Assistant", TEAL, "Alert triage for Visual Investigator"),
    ("Procurement Integrity Analyst", GREEN, "A watchdog over every tender"),
    ("Global Intelligence", AMBER, "The world's playbook — with sources"),
]
ty = 1.9
for i, (t, c, d) in enumerate(team):
    y = ty + i * 1.06
    rect(s, MX, y + 0.06, 0.14, 0.14, c)
    tf = box(s, MX + 0.34, y, 4.6, 0.35)
    run(para(tf, True), t, 15.5, NAVY, bold=True)
    tf = box(s, MX + 0.34, y + 0.36, 4.6, 0.4)
    run(para(tf, True), d, 12, GRAY)
tf = box(s, MX, ty + 4.45, 4.7, 0.4)
p = para(tf, True)
run(p, "One window. Pick an agent. Ask.", 14, BLUE, bold=True)
browser(s, 5.6, 1.62, 7.1, "shot_dropdown.jpg")
footer(s, 3)

# ════════════════════════════════════════════════════════════════════
# 4 · AGENT 01 — SAS VIYA COPILOT
# ════════════════════════════════════════════════════════════════════
s = slide()
header(s, "Agent 01 · SAS Viya Copilot", "Ask. It Finds, Builds, Explains.",
       accent=BLUE)
side_bullets(s, [
    "Finds and profiles your data",
    "Builds models with SAS AutoML",
    "Scores new cases in real time",
    "Answers from official SAS docs",
], BLUE)
tf = box(s, MX, 5.6, 4.35, 0.9)
p = para(tf, True, line=1.25)
run(p, "Behind the scenes: ", 12.5, NAVY, bold=True)
run(p, "a team of six specialist agents — visible in the trace.", 12.5, GRAY)
browser(s, 5.35, 1.62, 7.35, "shot_copilot.jpg")
footer(s, 4)

# ════════════════════════════════════════════════════════════════════
# 5 · DASHBOARDS ON DEMAND
# ════════════════════════════════════════════════════════════════════
s = slide()
header(s, "Agent 01 · Visual Analytics", "Dashboards on Demand", accent=BLUE)
ty = 2.0
for i, wrd in enumerate(["Show it.", "Analyze it.", "Create it."]):
    tf = box(s, MX, ty + i * 0.72, 4.3, 0.65)
    run(para(tf, True), wrd, 30, (NAVY, BLUE, NAVY)[i % 3], bold=True)
tf = box(s, MX, ty + 2.55, 4.3, 1.6)
p = para(tf, True, line=1.3)
run(p, "It advises like a BI consultant:", 12.5, GRAY)
p = para(tf, line=1.3, space_before=4)
run(p, "“Add a KPI for single-bid awards, so you can track competition "
       "at a glance.”", 13, NAVY, italic=True)
browser(s, 5.35, 1.62, 7.35, "shot_dashboard.jpg")
footer(s, 5)

# ════════════════════════════════════════════════════════════════════
# 6 · AGENT 02 — INVESTIGATION ASSISTANT
# ════════════════════════════════════════════════════════════════════
s = slide()
header(s, "Agent 02 · SAS Visual Investigator", "Your Morning Briefing.",
       accent=TEAL)
side_bullets(s, [
    "Prioritizes the alert queue",
    "Explains why each alert fired",
    "Maps hidden networks",
    "Flags likely false positives",
], TEAL)
tf = box(s, MX, 5.6, 4.35, 0.9)
p = para(tf, True, line=1.25)
run(p, "Investigators investigate — ", 12.5, NAVY, bold=True)
run(p, "the sorting is already done.", 12.5, GRAY)
browser(s, 5.35, 1.62, 7.35, "shot_vi.jpg")
footer(s, 6)

# ════════════════════════════════════════════════════════════════════
# 7 · AGENT 03 — PROCUREMENT INTEGRITY
# ════════════════════════════════════════════════════════════════════
s = slide()
header(s, "Agent 03 · Procurement Integrity", "Red Flags Before Award.",
       accent=GREEN)
side_bullets(s, [
    "Supplier risk scores",
    "Bid-rigging screen",
    "Price-anomaly detection",
    "Split purchases & duplicate invoices",
], GREEN)
tf = box(s, MX, 5.6, 4.35, 0.9)
p = para(tf, True, line=1.25)
run(p, "Catch it before contract award — ", 12.5, NAVY, bold=True)
run(p, "when it's cheapest to fix.", 12.5, GRAY)
browser(s, 5.35, 1.62, 7.35, "shot_procurement.jpg")
footer(s, 7)

# ════════════════════════════════════════════════════════════════════
# 8 · AGENT 04 — GLOBAL INTELLIGENCE
# ════════════════════════════════════════════════════════════════════
s = slide()
header(s, "Agent 04 · Global Intelligence", "The World, Cited.", accent=AMBER)
side_bullets(s, [
    "What other countries are doing",
    "Emerging tech, summarized",
    "Monitored topics, on schedule",
    "Every answer carries its sources",
], AMBER)
tf = box(s, MX, 5.6, 4.35, 0.9)
p = para(tf, True, line=1.25)
run(p, "Benchmark decisions — ", 12.5, NAVY, bold=True)
run(p, "without standing up a research team.", 12.5, GRAY)
browser(s, 5.35, 1.62, 7.35, "shot_web.jpg")
footer(s, 8)

# ════════════════════════════════════════════════════════════════════
# 9 · BUILT FOR THE KINGDOM — Arabic mode, real
# ════════════════════════════════════════════════════════════════════
s = slide()
header(s, "Sovereignty & Trust", "Built for the Kingdom.")
browser(s, MX, 1.62, 7.35, "shot_arabic.jpg")
tx = MX + 7.75
tiles = [
    ("عربي — one click", "The whole experience flips: interface, "
                         "answers, voice."),
    ("Data stays home", "Runs with an on-premises AI model. Nothing "
                        "leaves."),
    ("Every answer shows its work", "Steps, tools, and data — "
                                    "audit-ready."),
]
for i, (t, d) in enumerate(tiles):
    y = 1.75 + i * 1.55
    rect(s, tx, y, W - MX - tx, 1.35, TINT, line_color=BORDER, round_=True)
    rect(s, tx, y + 0.12, 0.055, 1.11, BLUE)
    tf = box(s, tx + 0.28, y + 0.18, W - MX - tx - 0.5, 0.45)
    run(para(tf, True, line=1.05), t, 14.5, NAVY, bold=True)
    tf = box(s, tx + 0.28, y + 0.68, W - MX - tx - 0.5, 0.6)
    run(para(tf, True, line=1.15), d, 11.5, INK)
footer(s, 9)

# ════════════════════════════════════════════════════════════════════
# 10 · LIVE DEMO — beats + the trace
# ════════════════════════════════════════════════════════════════════
s = slide()
header(s, "Live Demo", "What You'll See Today")
demo = [
    (BLUE,  "A model built from one sentence"),
    (BLUE,  "A dashboard, shown and analyzed in chat"),
    (TEAL,  "Today's alerts, triaged"),
    (GREEN, "A bid-rotation ring, uncovered"),
    (AMBER, "The world scanned — with sources"),
]
for i, (c, txt) in enumerate(demo):
    y = 1.8 + i * 0.92
    sp = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(MX), Inches(y),
                            Inches(0.52), Inches(0.52))
    sp.fill.solid(); sp.fill.fore_color.rgb = RGBColor.from_string(c)
    sp.line.fill.background(); sp.shadow.inherit = False
    tf = sp.text_frame
    tf.word_wrap = False
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, 0)
    run(para(tf, True, align=PP_ALIGN.CENTER), str(i + 1), 16, WHITE,
        bold=True)
    tf = box(s, MX + 0.76, y, 5.7, 0.56, anchor=MSO_ANCHOR.MIDDLE)
    run(para(tf, True, line=1.1), txt, 15.5, INK)
browser(s, 7.15, 1.95, 5.55, "shot_trace.jpg")
tf = box(s, 7.15, 5.85, 5.55, 0.4)
p = para(tf, True, align=PP_ALIGN.CENTER)
run(p, "Every step the agents take — visible.", 12.5, NAVY, bold=True)
footer(s, 10)

# ════════════════════════════════════════════════════════════════════
# 11 · NEXT STEPS (navy)
# ════════════════════════════════════════════════════════════════════
s = slide()
rect(s, 0, 0, W, H, NAVY)
rect(s, 0, 0, W, 0.07, BLUE)
rect(s, MX, 0.62, 0.42, 0.045, SKY)
tf = box(s, MX + 0.56, 0.50, CW - 0.56, 0.3)
run(para(tf, True), "NEXT STEPS", 11, SKY, bold=True)
tf = box(s, MX, 0.84, CW, 0.62)
run(para(tf, True), "From Demo to Production", 27, WHITE, bold=True)
steps = [
    ("1", "Pick the use cases", "Procurement integrity is ready today."),
    ("2", "Connect your platform", "Your SAS, your data, your AI model."),
    ("3", "Pilot & scale", "Measure. Expand agent by agent."),
]
cw, gapx = 3.86, 0.26
for i, (num, t, d) in enumerate(steps):
    x = MX + i * (cw + gapx)
    rect(s, x, 2.1, cw, 2.5, NAVY_2, round_=True)
    tf = box(s, x + 0.32, 2.4, cw - 0.64, 0.5)
    run(para(tf, True), num, 24, SKY, bold=True)
    tf = box(s, x + 0.32, 2.98, cw - 0.64, 0.55)
    run(para(tf, True, line=1.05), t, 16.5, WHITE, bold=True)
    tf = box(s, x + 0.32, 3.56, cw - 0.64, 0.8)
    run(para(tf, True, line=1.2), d, 12, "C9DAEE")
tf = box(s, MX, 5.5, CW, 0.9)
p = para(tf, True, align=PP_ALIGN.CENTER)
run(p, "شكراً", 30, WHITE, bold=True)
run(p, "    —    Thank You", 30, WHITE, bold=True)
tf = box(s, MX, 6.42, CW, 0.4)
run(para(tf, True, align=PP_ALIGN.CENTER),
    "Raed Aldweik  ·  SAS Middle East", 12.5, SKY)
footer(s, 11, dark=True)

prs.core_properties.title = "The NCGR AI Copilot — SAS Agentic AI in Action"
prs.core_properties.author = "SAS Middle East"
OUT = str(HERE / "NCGR_AI_Copilot_v2.pptx")
prs.save(OUT)
print("saved", OUT)
