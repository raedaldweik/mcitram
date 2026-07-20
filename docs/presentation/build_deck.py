# Builds the NCGR meeting deck in the SAS corporate presentation style.
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ── palette (SAS corporate) ─────────────────────────────────────────
NAVY      = "032954"   # SAS midnight blue
NAVY_2    = "0B3A6B"   # lighter navy for cards on navy
BLUE      = "0766D1"   # SAS blue
SKY       = "8FC1F2"   # light blue for text on navy
TINT      = "EDF4FC"   # light blue card fill
BORDER    = "D7E4F4"
INK       = "1F2A37"   # body text
GRAY      = "53565A"   # SAS gray
FAINT     = "9AA5B1"   # footers
WHITE     = "FFFFFF"

TEAL      = "0E7490"
GREEN     = "0B6E4F"
AMBER     = "B45309"

from pathlib import Path
LOGO = str(Path(__file__).resolve().parents[2] / "frontend" / "public" / "ncgr-logo.png")

W, H = 13.333, 7.5
MX = 0.62                      # side margin
CW = W - 2 * MX                # content width

prs = Presentation()
prs.slide_width = Inches(W)
prs.slide_height = Inches(H)
BLANK = prs.slide_layouts[6]

FOOT_L = "Company Confidential — For Internal Use Only"
FOOT_C = "Copyright © SAS Institute Inc. All rights reserved."


# ── helpers ─────────────────────────────────────────────────────────
def slide():
    return prs.slides.add_slide(BLANK)

def box(s, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, 0)
    return tf

def para(tf, first=False, align=PP_ALIGN.LEFT, space_after=0, space_before=0,
         line=None):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_after = Pt(space_after)
    p.space_before = Pt(space_before)
    if line:
        p.line_spacing = line
    return p

def run(p, text, size, color, bold=False, italic=False, name="Arial"):
    r = p.add_run()
    r.text = text
    f = r.font
    f.size, f.bold, f.italic, f.name = Pt(size), bold, italic, name
    f.color.rgb = RGBColor.from_string(color)
    return r

def rect(s, x, y, w, h, fill, line_color=None, round_=False, adj=0.055):
    shape_t = MSO_SHAPE.ROUNDED_RECTANGLE if round_ else MSO_SHAPE.RECTANGLE
    sp = s.shapes.add_shape(shape_t, Inches(x), Inches(y), Inches(w), Inches(h))
    if round_:
        try:
            sp.adjustments[0] = adj
        except Exception:
            pass
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = RGBColor.from_string(fill)
    if line_color is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = RGBColor.from_string(line_color)
        sp.line.width = Pt(1)
    sp.shadow.inherit = False
    return sp

def footer(s, n, dark=False):
    c = "5E7A9B" if dark else FAINT
    tf = box(s, MX, H - 0.34, 5.6, 0.25)
    run(para(tf, True), FOOT_L, 7.5, c)
    tf = box(s, W / 2 - 2.4, H - 0.34, 4.8, 0.25)
    run(para(tf, True, align=PP_ALIGN.CENTER), FOOT_C, 7.5, c)
    tf = box(s, W - MX - 2.0, H - 0.34, 2.0, 0.25)
    run(para(tf, True, align=PP_ALIGN.RIGHT), f"sas.com   ·   {n}", 7.5, c)

def header(s, kicker, title, accent=BLUE, title_color=NAVY):
    rect(s, MX, 0.52, 0.42, 0.045, accent)             # accent dash
    tf = box(s, MX + 0.56, 0.40, CW - 0.56, 0.3)
    run(para(tf, True), kicker.upper(), 11, accent, bold=True)
    tf = box(s, MX, 0.74, CW, 0.62)
    run(para(tf, True), title, 27, title_color, bold=True)

def bullets(tf, items, size=13.5, color=INK, accent=BLUE, gap=9, first=True,
            line=1.12):
    for i, it in enumerate(items):
        p = para(tf, first and i == 0, space_after=gap, line=line)
        run(p, "▪  ", size, accent, bold=True)
        if isinstance(it, tuple):                       # (bold lead, rest)
            run(p, it[0], size, color, bold=True)
            run(p, it[1], size, color)
        else:
            run(p, it, size, color)

def ask_card(s, x, y, w, h, questions, accent=SKY):
    rect(s, x, y, w, h, NAVY, round_=True, adj=0.045)
    tf = box(s, x + 0.3, y + 0.26, w - 0.6, 0.3)
    run(para(tf, True), "JUST ASK", 10.5, accent, bold=True)
    tf = box(s, x + 0.3, y + 0.62, w - 0.6, h - 0.9)
    for i, q in enumerate(questions):
        p = para(tf, i == 0, space_after=10, line=1.15)
        run(p, f"“{q}”", 12.5, WHITE, italic=True)

def value_strip(s, text, y=6.32, accent=BLUE):
    rect(s, MX, y, CW, 0.62, TINT, round_=True, adj=0.14)
    rect(s, MX, y + 0.1, 0.055, 0.42, accent)
    tf = box(s, MX + 0.28, y, CW - 0.56, 0.62, anchor=MSO_ANCHOR.MIDDLE)
    p = para(tf, True, line=1.1)
    run(p, "Why it matters — ", 12.5, NAVY, bold=True)
    run(p, text, 12.5, INK)

def demo_card(s, x, y, w, text, accent):
    rect(s, x, y, w, 0.8, WHITE, line_color=BORDER, round_=True, adj=0.12)
    rect(s, x, y + 0.12, 0.055, 0.56, accent)
    tf = box(s, x + 0.26, y, w - 0.5, 0.8, anchor=MSO_ANCHOR.MIDDLE)
    p = para(tf, True, line=1.12)
    run(p, "IN THE DEMO — ", 11, accent, bold=True)
    run(p, text, 11.5, INK)

def chip(s, x, y, text, color):
    w = 0.10 * len(text) + 0.32
    rect(s, x, y, w, 0.34, TINT, line_color=BORDER, round_=True, adj=0.5)
    tf = box(s, x, y, w, 0.34, anchor=MSO_ANCHOR.MIDDLE)
    run(para(tf, True, align=PP_ALIGN.CENTER), text, 10.5, color, bold=True)
    return w


# ════════════════════════════════════════════════════════════════════
# 1 · TITLE (navy)
# ════════════════════════════════════════════════════════════════════
s = slide()
rect(s, 0, 0, W, H, NAVY)
rect(s, 0, 0, W, 0.07, BLUE)
s.shapes.add_picture(LOGO, Inches(W - MX - 1.05), Inches(0.55),
                     height=Inches(1.05))

rect(s, MX, 2.02, 0.6, 0.055, BLUE)
tf = box(s, MX, 2.28, 11.2, 0.35)
run(para(tf, True), "SAS MIDDLE EAST  ·  AGENTIC AI IN ACTION", 13, SKY,
    bold=True)

tf = box(s, MX, 2.72, 12.1, 1.05)
run(para(tf, True), "The NCGR AI Copilot", 44, WHITE, bold=True)

tf = box(s, MX, 3.82, 12.1, 0.55)
p = para(tf, True)
run(p, "«مُعين»", 22, SKY, bold=True)
run(p, "   —   One assistant. Four specialist agents. Your platform, "
       "your data, your choice of AI model.", 17, SKY)

tf = box(s, MX, 6.28, 9.0, 0.75)
p = para(tf, True, space_after=3)
run(p, "National Center for Government Resources Systems (NCGR)", 13, WHITE,
    bold=True)
p = para(tf)
run(p, "SAS Middle East  ·  July 2026", 12, "AFC2DA")
footer(s, 1, dark=True)

# ════════════════════════════════════════════════════════════════════
# 2 · THE VISION
# ════════════════════════════════════════════════════════════════════
s = slide()
header(s, "The Vision", "Your Copilot, On Your Terms")
tf = box(s, MX, 1.42, CW, 0.42)
run(para(tf, True),
    "A conversational assistant for NCGR — the convenience of a copilot, "
    "without giving up control of your data or your AI.", 14.5, GRAY)

cards = [
    ("1", "Your choice of AI model",
     "Works with the AI model you choose — in the cloud today, fully "
     "on-premises tomorrow. Swap it at any time; nothing else changes."),
    ("2", "The full power of SAS",
     "Every answer is produced by your governed SAS platform — real "
     "data, real models, real dashboards. Not a chatbot’s opinion."),
    ("3", "Trust you can inspect",
     "Every answer shows the steps and data behind it — and the whole "
     "experience switches between English and Arabic at one click."),
]
cw, gapx = 3.86, 0.26
for i, (num, t, d) in enumerate(cards):
    x = MX + i * (cw + gapx)
    rect(s, x, 2.15, cw, 3.55, TINT, line_color=BORDER, round_=True)
    rect(s, x, 2.15, cw, 0.09, BLUE)
    tf = box(s, x + 0.32, 2.55, cw - 0.64, 0.55)
    run(para(tf, True), num, 26, BLUE, bold=True)
    tf = box(s, x + 0.32, 3.18, cw - 0.64, 0.65)
    run(para(tf, True, line=1.05), t, 17, NAVY, bold=True)
    tf = box(s, x + 0.32, 3.86, cw - 0.64, 1.7)
    run(para(tf, True, line=1.22), d, 12.5, INK)

value_strip(s, "the copilot experience your teams want, on the "
               "sovereignty terms the Kingdom requires.", y=6.1)
footer(s, 2)

# ════════════════════════════════════════════════════════════════════
# 3 · MEET YOUR AI TEAM
# ════════════════════════════════════════════════════════════════════
s = slide()
header(s, "One Chat · Four Agents", "Meet Your AI Team")

team = [
    ("SAS Viya Copilot", BLUE,
     "Your front door to the SAS platform — data, models, and "
     "dashboards on request."),
    ("Investigation Assistant", TEAL,
     "Triage for SAS Visual Investigator — priorities, context, and "
     "next steps for every alert."),
    ("Procurement Integrity Analyst", GREEN,
     "A watchdog over tenders, bids, suppliers, and invoices — with "
     "ready risk models."),
    ("Global Intelligence", AMBER,
     "What the world is doing — news, benchmarks, and trends, always "
     "with sources."),
]
cw, gapx = 2.92, 0.19
for i, (t, c, d) in enumerate(team):
    x = MX + i * (cw + gapx)
    rect(s, x, 1.75, cw, 3.4, WHITE, line_color=BORDER, round_=True)
    rect(s, x, 1.75, cw, 0.09, c)
    tf = box(s, x + 0.26, 2.05, cw - 0.52, 0.32)
    run(para(tf, True), f"AGENT 0{i+1}", 10, c, bold=True)
    tf = box(s, x + 0.26, 2.38, cw - 0.52, 0.95)
    run(para(tf, True, line=1.05), t, 15.5, NAVY, bold=True)
    tf = box(s, x + 0.26, 3.35, cw - 0.52, 1.6)
    run(para(tf, True, line=1.22), d, 11.5, INK)

rect(s, MX, 5.45, CW, 0.72, NAVY, round_=True, adj=0.12)
tf = box(s, MX + 0.32, 5.45, CW - 0.64, 0.72, anchor=MSO_ANCHOR.MIDDLE)
p = para(tf, True, line=1.15)
run(p, "One chat window — pick the agent from a dropdown.  ", 12.5,
    WHITE, bold=True)
run(p, "Every agent works in English or Arabic, shows its reasoning trail, "
       "and answers with charts — not just text.", 12.5, SKY)
value_strip(s, "four everyday workflows covered on day one — and the "
               "same pattern extends to any department.", y=6.42)
footer(s, 3)

# ════════════════════════════════════════════════════════════════════
# 4 · AGENT 01 — SAS VIYA COPILOT
# ════════════════════════════════════════════════════════════════════
s = slide()
header(s, "Agent 01 · SAS Viya Copilot",
       "From Question to Model — in Minutes", accent=BLUE)

tf = box(s, MX, 1.62, 7.1, 0.35)
run(para(tf, True), "WHAT IT DOES", 11, GRAY, bold=True)
tf = box(s, MX, 1.98, 7.1, 2.7)
bullets(tf, [
    "Ask in plain language — it finds the right data, profiles it, and "
    "explains what’s in it.",
    "Builds and compares machine-learning models with SAS AutoML — then "
    "scores new cases in real time.",
    "Creates realistic demo data on request — ready for workshops and "
    "training.",
    "Answers “how do I…?” questions from official SAS "
    "documentation, with references.",
], accent=BLUE)

tf = box(s, MX, 4.62, 7.1, 0.32)
run(para(tf, True), "BEHIND THE SCENES, IT LEADS A TEAM OF SPECIALISTS", 10.5,
    GRAY, bold=True)
x = MX
y = 5.0
for name in ["Data Steward", "Data Engineer", "Model Builder",
             "Insights & Reporting", "Platform Guide", "Dashboard Designer"]:
    w_next = 0.10 * len(name) + 0.32
    if x + w_next > MX + 7.15:
        x = MX
        y += 0.46
    w = chip(s, x, y, name, NAVY)
    x += w + 0.14

ask_card(s, 8.0, 1.62, 4.7, 3.3, [
    "What data do we have about suppliers?",
    "Build a model that predicts payment delays — and score this "
    "invoice with it.",
    "How do I publish a model? Show me the official steps.",
])
demo_card(s, 8.0, 5.08, 4.7,
          "data explored and a model built, from one request.", BLUE)
value_strip(s, "analytics without the queue — business teams get "
               "answers in minutes, and everything stays inside your "
               "governed SAS environment.")
footer(s, 4)

# ════════════════════════════════════════════════════════════════════
# 5 · AGENT 01 (CONT.) — DASHBOARDS ON DEMAND
# ════════════════════════════════════════════════════════════════════
s = slide()
header(s, "Agent 01 · SAS Viya Copilot · Visual Analytics",
       "Dashboards on Demand", accent=BLUE)

half = (CW - 0.3) / 2
for i, (t, items) in enumerate([
    ("Show & analyze", [
        "“Show me the procurement dashboard and analyze it.”",
        "The dashboard appears right in the chat — together with the "
        "story behind the numbers.",
        "One click opens the live version in SAS Visual Analytics.",
    ]),
    ("Create & improve", [
        "Starts from your NCGR-branded template — your design, your "
        "identity.",
        "The copilot fills it with the data you name and saves it as a new "
        "dashboard, shown in the chat.",
        "It advises like a BI consultant: what to add, what to track, what "
        "to drop.",
    ]),
]):
    x = MX + i * (half + 0.3)
    rect(s, x, 1.72, half, 3.7, TINT if i == 0 else WHITE,
         line_color=BORDER, round_=True)
    rect(s, x, 1.72, half, 0.09, BLUE)
    tf = box(s, x + 0.34, 2.05, half - 0.68, 0.5)
    run(para(tf, True), t, 17, NAVY, bold=True)
    tf = box(s, x + 0.34, 2.68, half - 0.68, 3.0)
    bullets(tf, items, size=12.5, gap=10, accent=BLUE)

value_strip(s, "the copilot doesn’t just build what you ask — it "
               "recommends more:  “add a KPI for single-bid awards, so "
               "you can track competition at a glance.”")
footer(s, 5)

# ════════════════════════════════════════════════════════════════════
# 6 · AGENT 02 — INVESTIGATION ASSISTANT
# ════════════════════════════════════════════════════════════════════
s = slide()
header(s, "Agent 02 · SAS Visual Investigator",
       "A Morning Briefing for Every Investigator", accent=TEAL)

tf = box(s, MX, 1.62, 7.1, 0.35)
run(para(tf, True), "WHAT IT DOES", 11, GRAY, bold=True)
tf = box(s, MX, 1.98, 7.1, 3.4)
bullets(tf, [
    "Reads the alert queue and tells each investigator what to look at "
    "first — and why.",
    "Explains in plain language why an alert fired, and gathers the full "
    "picture around it.",
    "Maps the network: which people, companies, and accounts are connected "
    "— and how.",
    "Flags likely false positives and recommends the next action, so effort "
    "goes where it matters.",
], accent=TEAL, size=14.5, gap=14)

ask_card(s, 8.0, 1.62, 4.7, 3.3, [
    "What should I look at first today?",
    "Why did this alert fire — could it be a false positive?",
    "Who is connected to this supplier, and through what?",
])
demo_card(s, 8.0, 5.08, 4.7,
          "a briefing on today’s alert queue, then one case end-to-end.",
          TEAL)
value_strip(s, "investigators spend their day investigating — not "
               "sorting alerts — and new team members work like "
               "veterans from day one.", accent=TEAL)
footer(s, 6)

# ════════════════════════════════════════════════════════════════════
# 7 · AGENT 03 — PROCUREMENT INTEGRITY ANALYST
# ════════════════════════════════════════════════════════════════════
s = slide()
header(s, "Agent 03 · Procurement Integrity",
       "A Watchdog Over Every Tender", accent=GREEN)

tf = box(s, MX, 1.62, 7.1, 0.35)
run(para(tf, True), "WHAT IT DOES", 11, GRAY, bold=True)
tf = box(s, MX, 1.98, 7.1, 3.4)
bullets(tf, [
    "Watches tenders, bids, suppliers, and invoices across government "
    "entities for red flags.",
    ("Ready-made analytics: ", "supplier risk scores, a bid-rigging screen, "
     "and price-anomaly detection."),
    "Surfaces patterns people miss — rotating bid winners, purchases "
    "split to stay under approval thresholds, duplicate invoices.",
    "Answers with numbers and charts, not opinions — every figure "
    "drill-down-able in conversation.",
], accent=GREEN, size=14.5, gap=14)

ask_card(s, 8.0, 1.62, 4.7, 3.3, [
    "Who are our riskiest suppliers right now?",
    "Run the bid-rigging screen on IT tenders.",
    "How much are we overpaying versus market prices?",
])
demo_card(s, 8.0, 5.08, 4.7,
          "a bid-rotation ring uncovered — with the charts to prove it.",
          GREEN)
value_strip(s, "integrity issues surface before contract award — when "
               "they are cheapest to fix.", accent=GREEN)
footer(s, 7)

# ════════════════════════════════════════════════════════════════════
# 8 · AGENT 04 — GLOBAL INTELLIGENCE
# ════════════════════════════════════════════════════════════════════
s = slide()
header(s, "Agent 04 · Global Intelligence", "Eyes on the World",
       accent=AMBER)

tf = box(s, MX, 1.62, 7.1, 0.35)
run(para(tf, True), "WHAT IT DOES", 11, GRAY, bold=True)
tf = box(s, MX, 1.98, 7.1, 3.4)
bullets(tf, [
    "Scans global news and publications on demand — in seconds, not "
    "days.",
    "Benchmarks NCGR against the world: what other governments are doing on "
    "procurement oversight and AI.",
    "Tracks emerging technologies and summarizes what they mean for your "
    "roadmap.",
    "Every answer carries its sources — open the original articles in "
    "one click.",
], accent=AMBER, size=14.5, gap=14)

ask_card(s, 8.0, 1.62, 4.7, 3.3, [
    "What are other countries doing on AI-driven procurement oversight?",
    "What changed in agentic AI this month?",
    "Summarize best practice for supplier risk monitoring.",
])
demo_card(s, 8.0, 5.08, 4.7,
          "a live, cited scan of what other governments are doing.", AMBER)
value_strip(s, "decisions informed by what’s happening outside — "
               "without standing up a research team.", accent=AMBER)
footer(s, 8)

# ════════════════════════════════════════════════════════════════════
# 9 · BUILT FOR THE KINGDOM
# ════════════════════════════════════════════════════════════════════
s = slide()
header(s, "Sovereignty & Trust", "Built for the Kingdom")

tiles = [
    ("Arabic at one click",
     "The عربي button flips the whole experience — "
     "interface and answers — into Arabic, and back again. The choice "
     "is remembered."),
    ("Your data stays home",
     "Designed to run with an on-premises AI model: questions, data, and "
     "answers never have to leave your environment."),
    ("Every answer shows its work",
     "A full trail of the steps, tools, and data behind each answer — "
     "ready for audit and review at any time."),
    ("Grows with you",
     "Start with four agents; add one per department or use case. The same "
     "design carries forward into SAS RAM as it evolves."),
]
tw, th, gx, gy = (CW - 0.3) / 2, 1.98, 0.3, 0.26
for i, (t, d) in enumerate(tiles):
    x = MX + (i % 2) * (tw + gx)
    y = 1.72 + (i // 2) * (th + gy)
    rect(s, x, y, tw, th, TINT, line_color=BORDER, round_=True)
    rect(s, x, y + 0.14, 0.055, th - 0.28, BLUE)
    tf = box(s, x + 0.34, y + 0.24, tw - 0.68, 0.45)
    run(para(tf, True), t, 15.5, NAVY, bold=True)
    tf = box(s, x + 0.34, y + 0.74, tw - 0.68, th - 0.9)
    run(para(tf, True, line=1.22), d, 12, INK)

value_strip(s, "a copilot aligned with Vision 2030 digital-sovereignty "
               "goals — modern AI, on national terms.")
footer(s, 9)

# ════════════════════════════════════════════════════════════════════
# 10 · WHAT YOU'LL SEE TODAY
# ════════════════════════════════════════════════════════════════════
s = slide()
header(s, "Live Demo", "What You’ll See Today")

demo = [
    (BLUE,  "The SAS Viya Copilot explores data and builds a model — "
            "from a single plain-language request."),
    (BLUE,  "A dashboard appears in the chat, analyzed — then a new one "
            "is created from the NCGR template."),
    (TEAL,  "The Investigation Assistant briefs us on today’s alert "
            "queue and digs into one case."),
    (GREEN, "The Procurement Analyst uncovers a bid-rotation ring — "
            "with the charts to prove it."),
    (AMBER, "Global Intelligence reports what other countries are doing "
            "— with cited sources."),
]
for i, (c, txt) in enumerate(demo):
    y = 1.72 + i * 0.86
    sp = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(MX), Inches(y),
                            Inches(0.52), Inches(0.52))
    sp.fill.solid(); sp.fill.fore_color.rgb = RGBColor.from_string(c)
    sp.line.fill.background(); sp.shadow.inherit = False
    tf = sp.text_frame
    tf.word_wrap = False
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, 0)
    run(para(tf, True, align=PP_ALIGN.CENTER), str(i + 1), 16, WHITE,
        bold=True)
    tf = box(s, MX + 0.78, y, CW - 0.78, 0.6, anchor=MSO_ANCHOR.MIDDLE)
    run(para(tf, True, line=1.1), txt, 14.5, INK)

tf = box(s, MX, 6.22, CW, 0.4)
p = para(tf, True)
run(p, "All live — ", 12.5, NAVY, bold=True)
run(p, "every step the agents take is visible on screen as they work.",
    12.5, GRAY)
footer(s, 10)

# ════════════════════════════════════════════════════════════════════
# 11 · NEXT STEPS (navy)
# ════════════════════════════════════════════════════════════════════
s = slide()
rect(s, 0, 0, W, H, NAVY)
rect(s, 0, 0, W, 0.07, BLUE)
rect(s, MX, 0.62, 0.42, 0.045, SKY)
tf = box(s, MX + 0.56, 0.50, CW - 0.56, 0.3)
run(para(tf, True), "NEXT STEPS", 11, SKY, bold=True)
tf = box(s, MX, 0.84, CW, 0.62)
run(para(tf, True), "From Demo to Production", 27, WHITE, bold=True)

steps = [
    ("1", "Choose the first use cases",
     "Pick two or three high-value workflows — procurement integrity "
     "is ready to start today."),
    ("2", "Connect your environment",
     "Your SAS platform, your data, and the AI model of your choice — "
     "on-premises or cloud."),
    ("3", "Pilot with your teams",
     "Measure time saved and decision quality — then expand, agent by "
     "agent."),
]
cw, gapx = 3.86, 0.26
for i, (num, t, d) in enumerate(steps):
    x = MX + i * (cw + gapx)
    rect(s, x, 1.95, cw, 2.9, NAVY_2, round_=True)
    tf = box(s, x + 0.32, 2.25, cw - 0.64, 0.5)
    run(para(tf, True), num, 24, SKY, bold=True)
    tf = box(s, x + 0.32, 2.83, cw - 0.64, 0.6)
    run(para(tf, True, line=1.05), t, 16, WHITE, bold=True)
    tf = box(s, x + 0.32, 3.48, cw - 0.64, 1.25)
    run(para(tf, True, line=1.22), d, 12, "C9DAEE")

tf = box(s, MX, 5.6, CW, 0.9)
p = para(tf, True, align=PP_ALIGN.CENTER)
run(p, "شكراً", 30, WHITE, bold=True)
run(p, "    —    Thank You", 30, WHITE, bold=True)
tf = box(s, MX, 6.5, CW, 0.4)
run(para(tf, True, align=PP_ALIGN.CENTER),
    "Raed Aldweik  ·  SAS Middle East", 12.5, SKY)
footer(s, 11, dark=True)

prs.core_properties.title = "The NCGR AI Copilot — SAS Agentic AI in Action"
prs.core_properties.author = "SAS Middle East"

OUT = str(Path(__file__).resolve().parent / "NCGR_AI_Copilot.pptx")
prs.save(OUT)
print("saved", OUT)
